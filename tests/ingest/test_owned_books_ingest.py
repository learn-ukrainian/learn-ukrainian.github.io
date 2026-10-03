"""Synthetic private fixtures prove extraction, full accounting and corpus rehearsal."""

import argparse
import io
import sqlite3
import zipfile
from pathlib import Path

import pymupdf
import pytest
import yaml
from docx import Document

from scripts.ingest import owned_books_ingest as owned

SCHEMA = """
CREATE TABLE textbooks (
 id INTEGER PRIMARY KEY, chunk_id TEXT UNIQUE, title TEXT, text TEXT,
 source_file TEXT, grade TEXT, author TEXT, author_uk TEXT, char_count INTEGER
);
CREATE VIRTUAL TABLE textbooks_fts USING fts5(title,text,content='textbooks',content_rowid='id');
CREATE TRIGGER textbooks_ai AFTER INSERT ON textbooks BEGIN
 INSERT INTO textbooks_fts(rowid,title,text) VALUES(new.id,new.title,new.text);
END;
"""


def pdf(*pages, encrypted=False):
    with pymupdf.open() as doc:
        for text in pages:
            page = doc.new_page()
            if text:
                page.insert_text((40, 40), text)
        return (
            doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="synthetic", user_pw="synthetic")
            if encrypted
            else doc.tobytes()
        )


def archive(members):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as z:
        for name, content in members:
            z.writestr(name, content)
    return stream.getvalue()


def docx(*paragraphs):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Synthetic table content."
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def epub():
    return archive(
        [
            (
                "META-INF/container.xml",
                '<container><rootfiles><rootfile full-path="book/package.opf"/></rootfiles></container>',
            ),
            (
                "book/package.opf",
                '<package><manifest><item id="z" href="z.xhtml"/><item id="a" href="a.xhtml"/></manifest><spine><itemref idref="z"/><itemref idref="a"/></spine></package>',
            ),
            ("book/a.xhtml", "<html><body>Second chapter.</body></html>"),
            ("book/z.xhtml", "<html><body>First chapter.<script>Hidden code.</script></body></html>"),
        ]
    )


def pptx():
    return archive(
        [
            (
                "ppt/presentation.xml",
                '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="second"/><p:sldId id="257" r:id="first"/></p:sldIdLst></p:presentation>',
            ),
            (
                "ppt/_rels/presentation.xml.rels",
                '<Relationships><Relationship Id="first" Target="slides/slide1.xml"/><Relationship Id="second" Target="slides/slide2.xml"/></Relationships>',
            ),
            (
                "ppt/slides/slide1.xml",
                '<slide xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:p><a:t>Last </a:t><a:t>slide.</a:t></a:p></slide>',
            ),
            (
                "ppt/slides/slide2.xml",
                '<slide xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:p><a:t>First </a:t><a:t>slide.</a:t></a:p></slide>',
            ),
        ]
    )


@pytest.fixture
def env(tmp_path, monkeypatch):
    rights = owned.load_owned_rights()
    rights["owned-synthetic-work"] = {"rights": "owned_cite_only"}
    monkeypatch.setattr(owned, "load_owned_rights", lambda: rights)
    db = tmp_path / "corpus.db"
    with sqlite3.connect(db) as conn:
        conn.executescript(SCHEMA)
    root = tmp_path / "drive-ukrainian"
    root.mkdir()
    inventory = tmp_path / "INVENTORY.yaml"
    return db, root, inventory, tmp_path / "private-jsonl"


def row(**updates):
    result = {
        "id": "synthetic-work",
        "title": "Synthetic reference",
        "author": "Synthetic author",
        "rights": "owned_cite_only",
        "files": ["**/*.pdf"],
    }
    result.update(updates)
    return result


def args(env, rows=None, **updates):
    db, _root, inventory, out = env
    inventory.write_text(yaml.safe_dump({"schema": 1, "works": rows or [row()]}))
    result = argparse.Namespace(
        inventory=inventory, out_dir=out, db=db, check=False, dry_run=False, force=False, only=None
    )
    for key, value in updates.items():
        setattr(result, key, value)
    return result


def test_ingest_noop_check_force_and_stale_input(env):
    db, root, _inventory, out = env
    file = root / "fixture.pdf"
    file.write_bytes(pdf("Extraordinaryword original sentence."))
    request = args(env)
    report, code = owned.run(request)
    assert code == 0 and report["rows"][0]["status"] == "ingested"
    assert report["rows"][0]["retrieval"]["chunk_id"].startswith("owned-synthetic-work_")
    before = db.read_bytes(), {p.name: (p.read_bytes(), p.stat().st_mtime_ns) for p in out.iterdir()}
    assert owned.run(request)[1] == 0
    assert before == (db.read_bytes(), {p.name: (p.read_bytes(), p.stat().st_mtime_ns) for p in out.iterdir()})
    request.check = True
    assert owned.run(request)[1] == 0
    file.write_bytes(pdf("Replacementword completely changed sentence."))
    assert owned.run(request)[0]["rows"][0]["status"] == "error:stale_input"
    request.check = False
    request.force = True
    assert owned.run(request)[1] == 0
    with sqlite3.connect(db) as conn:
        assert not conn.execute(
            "SELECT rowid FROM textbooks_fts WHERE textbooks_fts MATCH 'Extraordinaryword'"
        ).fetchall()
        assert conn.execute("SELECT rowid FROM textbooks_fts WHERE textbooks_fts MATCH 'Replacementword'").fetchall()
        assert conn.execute("SELECT count(*) FROM textbook_sections").fetchone()[0] == 1
    assert "Replacementword" in (out / "owned-synthetic-work.jsonl").read_text()


@pytest.mark.parametrize("mode", ["dry_run", "check"])
def test_read_only_modes_never_write(env, mode):
    db, root, _inventory, out = env
    (root / "fixture.pdf").write_bytes(pdf("Synthetic readable sentence."))
    request = args(env, **{mode: True})
    before = db.read_bytes()
    report, code = owned.run(request)
    assert code == (mode == "check")
    assert report["rows"][0]["status"] == ("error:not_ingested" if mode == "check" else "ingested")
    assert db.read_bytes() == before and not out.exists()


@pytest.mark.parametrize(
    "extension,content,expected",
    [
        (".pdf", lambda: pdf("First sentence.", "Second sentence."), ["First sentence.", "Second sentence."]),
        (
            ".docx",
            lambda: docx("First paragraph.", "Second paragraph."),
            ["First paragraph.", "Second paragraph.", "Synthetic table content."],
        ),
        (".epub", epub, ["First chapter.", "Second chapter."]),
        (".pptx", pptx, ["First slide.", "Last slide."]),
    ],
)
def test_format_extraction_follows_document_order(extension, content, expected):
    units, pages, status = owned.extract(content(), extension)
    text = "\n".join(t for _, t in units)
    assert status == "extracted"
    assert all(p["status"] == "text" for p in pages)
    assert [text.index(part) for part in expected] == sorted(text.index(part) for part in expected)
    assert "Hidden code" not in text


def test_mixed_and_scanned_pdf_account_for_every_page():
    units, pages, status = owned.extract(pdf("Synthetic content.", "", "More content."), ".pdf")
    assert status == "extracted" and len(units) == 2
    assert [p["status"] for p in pages] == ["text", "page_no_text", "text"]
    units, pages, status = owned.extract(pdf("Synthetic content.", "", ""), ".pdf")
    assert status == "skipped:scanned_needs_ocr" and not units and len(pages) == 3
    assert owned.extract(pdf(""), ".pdf")[2] == "skipped:scanned_needs_ocr"


@pytest.mark.parametrize(
    "payload,extension,code",
    [
        (lambda: pdf("Protected.", encrypted=True), ".pdf", "encrypted"),
        (lambda: b"not a PDF", ".pdf", "corrupt"),
        (lambda: b"not a docx", ".docx", "corrupt"),
        (lambda: archive([]), ".epub", "corrupt"),
        (lambda: archive([]), ".pptx", "corrupt"),
        (lambda: b"unknown content", ".unknown", "unsupported_format"),
    ],
)
def test_closed_extraction_errors(payload, extension, code):
    with pytest.raises(owned.IngestError, match=f"^{code}$"):
        owned.extract(payload(), extension)


def test_archive_sorting_nested_depth_ignored_members_and_size_cap(env, monkeypatch):
    _db, root, _inventory, out = env
    a = pdf("Alphabeticalword first content.")
    z = pdf("Zoologicalword second content.")
    packed = archive(
        [
            ("z.pdf", z),
            ("a.pdf", a),
            ("__MACOSX/secret.pdf", z),
            (".DS_Store", "secret"),
            ("._a.pdf", b"broken AppleDouble"),
            ("folder/._z.pdf", b"broken AppleDouble"),
            ("nested.zip", archive([("inner.pdf", z), ("deeper.zip", archive([("deep.pdf", z)]))])),
            ("audio.mp3", "sound"),
            ("bad.bin", b"unknown"),
        ]
    )
    (root / "fixture.zip").write_bytes(packed)
    request = args(env, [row(files=["*.zip"])], dry_run=True)
    report, code = owned.run(request)
    assert code == 1  # unknown leaf never hidden by another successful chunk
    file_reports = report["rows"][0]["files"]
    assert len(file_reports) == 8
    assert file_reports[1]["file_index"] == "1.1"
    assert any(f["status"] == "skipped:nested_archive_depth" for f in file_reports)
    assert any(f["status"] == "error:unsupported_format" for f in file_reports)
    assert any(f["status"] == "skipped:audio_or_deck" for f in file_reports)
    monkeypatch.setattr(owned, "MEMBER_CAP", 10)
    report, code = owned.run(request)
    assert any(f["status"] == "skipped:member_too_large" for f in report["rows"][0]["files"])
    assert not out.exists()


def test_pdf_only_archive_filter(env):
    _db, root, _inventory, _out = env
    (root / "fixture.zip").write_bytes(
        archive([("file.pdf", pdf("Syntheticword readable text.")), ("file.docx", docx("Ignore me."))])
    )
    report, code = owned.run(args(env, [row(files=["*.zip"], ingest="pdf_text_only")]))
    assert code == 0 and report["rows"][0]["chunks"] == 1
    assert report["rows"][0]["files"][-1]["status"] == "extracted"
    assert any(f["status"] == "skipped:pdf_text_only" for f in report["rows"][0]["files"])


@pytest.mark.parametrize("state", sorted(owned.SKIPS))
def test_typed_inventory_skip(env, state):
    report, code = owned.run(args(env, [row(ingest=state)], check=True))
    assert code == 0 and report["rows"][0]["status"] == f"skipped:{state}"


def test_legacy_identity_rights_only_and_missing_source(env, monkeypatch):
    monkeypatch.setattr(owned, "load_owned_rights", lambda: {"legacy-book": {"rights": "owned_cite_only"}})
    db, root, _inventory, out = env
    (root / "existing.pdf").write_bytes(pdf("Legacyword unchanged work."))
    request = args(env, [row(ingest="already_ingested", source_file="legacy-book", title=None, author=None)])
    assert owned.run(request)[0]["rows"][0]["status"] == "error:existing_source_missing"
    with sqlite3.connect(db) as conn:
        conn.execute(
            "INSERT INTO textbooks(chunk_id,source_file,text) VALUES('legacy-id','legacy-book','Existing body')"
        )
    before = db.read_bytes()
    assert owned.run(request)[0]["rows"][0]["status"] == "already_ingested"
    assert db.read_bytes() == before and not out.exists()


def test_missing_expected_file_and_unexplained_row_fail_check(env):
    request = args(env, check=True)
    report, code = owned.run(request)
    assert code == 1 and report["rows"][0]["status"] == "error:missing_file"
    assert report["rows"][0]["files"] == [{"file_index": "pattern-1", "status": "error:missing_file"}]
    request = args(env, [row(files=[])], check=True)
    with pytest.raises(owned.IngestError, match="invalid_inventory"):
        owned.run(request)


def test_partial_success_ingests_available_text_without_hiding_file_errors(env):
    db, root, _inventory, out = env
    (root / "good.pdf").write_bytes(pdf("Syntheticword readable text."))
    (root / "broken.pdf").write_bytes(b"corrupt")
    report, code = owned.run(args(env))
    assert code == 1 and report["rows"][0]["status"] == "error:file_failure"
    assert len(report["rows"][0]["files"]) == 2 and out.exists()
    with sqlite3.connect(db) as conn:
        assert conn.execute("select count(*) from textbooks").fetchone()[0] == 1
    request = args(env, check=True)
    assert owned.run(request)[1] == 1


def test_private_name_and_content_never_appear_in_diagnostics_or_db_metadata(env, capsys, caplog):
    db, root, inventory, out = env
    sentinel = "SENSITIVE_NAME_SENTINEL"
    (root / f"{sentinel}.docx").write_bytes(docx(f"{sentinel} Syntheticword content."))
    args(
        env,
        [row(id="teacher-a-slides", rights="private_permission", title=sentinel, author=sentinel, files=["*.docx"])],
    )
    assert owned.main(["--inventory", str(inventory), "--out-dir", str(out), "--db", str(db)]) == 0
    captured = capsys.readouterr()
    assert sentinel not in captured.out + captured.err + caplog.text
    with sqlite3.connect(db) as conn:
        metadata = conn.execute(
            "SELECT chunk_id,title,source_file,grade,author,author_uk,subject FROM textbooks"
        ).fetchall()
        sections = conn.execute("SELECT source_file,section_title,section_number FROM textbook_sections").fetchall()
    assert sentinel not in str(metadata) + str(sections)
    assert metadata[0][1] == "Teacher materials A" and metadata[0][4:6] == ("", "")
    assert sentinel not in Path("registry/sources/owned-rights.yaml").read_text()
    # Broken input also stays sanitized; messages never carry private filenames.
    (root / f"{sentinel}.docx").write_bytes(sentinel.encode())
    assert owned.main(["--inventory", str(inventory), "--out-dir", str(out), "--db", str(db), "--force"]) == 1
    captured = capsys.readouterr()
    assert sentinel not in captured.out + captured.err + caplog.text


def test_chunk_ids_and_metadata_are_stable_and_lossless(monkeypatch):
    monkeypatch.setattr(owned, "CHUNK_SIZE", 5)
    content = "abcdefghij whole content"
    result = owned.chunks(row(author=None), [("1", 3, content)])
    assert "".join(r["text"] for r in result) == content
    assert result == owned.chunks(row(author=None), [("1", 3, content)])
    assert result[0]["chunk_id"] == "owned-synthetic-work_f1_p0003_c0001"
    assert result[0]["author"] == result[0]["author_uk"] == ""


def test_saved_artifacts_and_db_links_must_match(env):
    db, root, _inventory, out = env
    (root / "fixture.pdf").write_bytes(pdf("Syntheticword test content."))
    request = args(env)
    assert owned.run(request)[1] == 0
    request.check = True
    with sqlite3.connect(db) as conn:
        conn.execute("UPDATE textbooks SET parent_section_id=NULL")
    assert owned.run(request)[0]["rows"][0]["status"] == "error:artifact_mismatch"
    request.check = False
    request.force = True
    assert owned.run(request)[1] == 0
    request.check = True
    (out / "owned-synthetic-work.jsonl").write_text("invalid json")
    assert owned.run(request)[0]["rows"][0]["status"] == "error:artifact_mismatch"
    (out / "owned-synthetic-work.manifest.json").write_text("invalid json")
    assert owned.run(request)[0]["rows"][0]["status"] == "error:invalid_receipt"


def test_path_refusals_and_cli_errors_are_sanitized(env, monkeypatch, capsys):
    db, _root, _inventory, out = env
    request = args(env)
    request.out_dir = owned.REPOSITORY / "private-output"
    with pytest.raises(owned.IngestError, match="output_inside_repository"):
        owned.run(request)
    request.out_dir = out
    monkeypatch.setattr(owned, "is_network_filesystem_path", lambda p: True)
    with pytest.raises(owned.IngestError, match="network_database_refused"):
        owned.run(request)
    monkeypatch.setattr(owned, "is_network_filesystem_path", lambda p: False)
    request.db = db.parent / "missing.db"
    with pytest.raises(owned.IngestError, match="database_missing"):
        owned.run(request)
    with pytest.raises(SystemExit) as exc:
        owned.main(["--SENSITIVE_NAME_SENTINEL"])
    assert exc.value.code == 2
    assert "SENSITIVE_NAME_SENTINEL" not in capsys.readouterr().err


def test_inventory_validation_unknown_id_and_entrypoint_failure(env, capsys):
    db, _root, inventory, out = env
    request = args(env, only="unregistered")
    with pytest.raises(owned.IngestError, match="unknown_id"):
        owned.run(request)
    inventory.write_text("invalid yaml: [")
    assert owned.main(["--inventory", str(inventory), "--out-dir", str(out), "--db", str(db)]) == 1
    assert "invalid_inventory" in capsys.readouterr().out


def test_unique_encrypted_zip_and_xml_entity_boundaries():
    info = zipfile.ZipInfo("hidden.pdf")
    info.file_size = owned.MEMBER_CAP + 1
    with pytest.raises(owned.IngestError, match="member_too_large"):
        owned.read_member(None, info)
    info.file_size = 3
    info.flag_bits = 1
    with pytest.raises(owned.IngestError, match="encrypted"):
        owned.read_member(None, info)
    tree = owned.xml(b'<!DOCTYPE p [<!ENTITY secret SYSTEM "file:///synthetic-secret">]><p>&secret;</p>')
    assert tree.text is None


def test_fts_delete_trigger_variant_and_retrieval_failure(env):
    db, _root, _inventory, _out = env
    with sqlite3.connect(db) as conn:
        conn.execute(
            "CREATE TRIGGER textbooks_ad AFTER DELETE ON textbooks BEGIN INSERT INTO textbooks_fts(textbooks_fts,rowid,title,text) VALUES('delete',old.id,old.title,old.text); END"
        )
        record = owned.chunks(row(), [("1", 1, "Syntheticword original content.")])
        owned.replace_work(conn, "owned-synthetic-work", record)
        record[0]["text"] = "Replacementword new content."
        owned.replace_work(conn, "owned-synthetic-work", record)
        assert not conn.execute("SELECT rowid FROM textbooks_fts WHERE textbooks_fts MATCH 'Syntheticword'").fetchall()
        assert owned.retrieval_proof(conn, "owned-synthetic-work")["token"]
        with pytest.raises(owned.IngestError, match="retrieval_unverified"):
            owned.retrieval_proof(conn, "unknown")


def test_rehearsal_jsonl_loads_separate_db_without_rewriting_artifacts(env):
    db, root, _inventory, out = env
    (root / "fixture.pdf").write_bytes(pdf("Syntheticword readable fixture."))
    request = args(env)
    assert owned.run(request)[1] == 0
    original = {p.name: (p.read_bytes(), p.stat().st_mtime_ns) for p in out.iterdir()}
    second = db.with_name("second.db")
    with sqlite3.connect(second) as conn:
        conn.executescript(SCHEMA)
    request.db = second
    request.check = True
    assert owned.run(request)[1] == 1
    request.check = False
    request.dry_run = True
    assert owned.run(request)[1] == 0
    with sqlite3.connect(second) as conn:
        assert conn.execute("SELECT count(*) FROM textbooks").fetchone()[0] == 0
    request.dry_run = False
    assert owned.run(request)[1] == 0
    assert original == {p.name: (p.read_bytes(), p.stat().st_mtime_ns) for p in out.iterdir()}
    request.check = True
    assert owned.run(request)[1] == 0


def test_unexplained_inventory_state_reports_that_row_and_keeps_denominator(env):
    report, code = owned.run(
        args(env, [row(ingest="unexplained"), row(id="skipped-work", ingest="missing_source")], check=True)
    )
    assert code == 1 and report["denominator"] == 2
    assert report["rows"][0]["status"] == "error:unknown_ingest"
    assert report["rows"][1]["status"] == "skipped:missing_source"


def test_pdf_page_error_does_not_hide_other_pages_or_private_exception(monkeypatch):
    original = pymupdf.Page.get_text

    def get_text(page, *args, **kwargs):
        if page.number == 1:
            raise ValueError("SENSITIVE_NAME_SENTINEL")
        return original(page, *args, **kwargs)

    monkeypatch.setattr(pymupdf.Page, "get_text", get_text)
    units, pages, status = owned.extract(pdf("First page.", "Second page.", "Third page."), ".pdf")
    assert status == "error:corrupt"
    assert [page["status"] for page in pages] == ["text", "error:corrupt", "text"]
    assert [page for page, _ in units] == [1, 3]
    assert "SENSITIVE_NAME_SENTINEL" not in str(pages)


def test_archive_crc_failure_accounts_for_leaf_and_continues_other_members(env):
    _db, root, _inventory, _out = env
    payload = pdf("Corruptibleword sentence.")
    data = bytearray(archive([("a.pdf", payload), ("b.pdf", pdf("Retainedword sentence."))]))
    offset = data.index(payload)
    data[offset] ^= 1
    (root / "archive.zip").write_bytes(data)
    report, code = owned.run(args(env, [row(files=["*.zip"])]))
    assert code == 1
    assert report["rows"][0]["files"][1]["status"] == "error:corrupt"
    assert report["rows"][0]["files"][2]["status"] == "extracted"
    assert report["rows"][0]["chunks"] == 1


def test_inline_markup_preserves_words_and_block_boundaries():
    payload = archive(
        [
            ("META-INF/container.xml", '<container><rootfile full-path="package.opf"/></container>'),
            (
                "package.opf",
                '<package><manifest><item id="c" href="chapter.xhtml"/></manifest><spine><itemref idref="c"/></spine></package>',
            ),
            ("chapter.xhtml", "<html><body><p>Extra<b>ordinary</b>word.</p><p>Next paragraph.</p></body></html>"),
        ]
    )
    assert owned.package_text(payload, ".epub") == [(1, "Extraordinaryword.\n\nNext paragraph.")]
    payload = archive(
        [
            ("ppt/presentation.xml", '<p xmlns:r="relationships"><sldId r:id="slide"/></p>'),
            ("ppt/_rels/presentation.xml.rels", '<rels><rel Id="slide" Target="slides/slide.xml"/></rels>'),
            (
                "ppt/slides/slide.xml",
                '<slide xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:p><a:t>Extra</a:t><a:t>ordinaryword.</a:t></a:p><a:p><a:t>Next paragraph.</a:t></a:p></slide>',
            ),
        ]
    )
    assert owned.package_text(payload, ".pptx") == [(1, "Extraordinaryword.\nNext paragraph.")]


@pytest.mark.parametrize("pattern", ["tree/**", "tree/**/*", "**", "**/*"])
def test_recursive_directory_glob_includes_direct_and_nested_files(env, pattern):
    _db, root, _inventory, _out = env
    (root / "tree/nested/deep").mkdir(parents=True)
    (root / "tree/empty").mkdir()
    direct = root / "tree/direct.pdf"
    nested = root / "tree/nested/deep/book.pdf"
    direct.write_bytes(pdf("Directword page."))
    nested.write_bytes(pdf("Nestedword page."))
    files, missing = owned.matched_files(row(files=[pattern]), root)
    assert files == sorted([direct, nested]) and missing == []


@pytest.mark.parametrize(
    "text,status",
    [
        ("ɭɤɪɚʀғɧɫɶɤɨɸ", "garbled_text_layer"),
        (owned.UKRAINIAN_LETTERS + owned.UKRAINIAN_LETTERS.lower(), "text"),
        ("Clean English text ABC xyz", "text"),
        ("abcdɭ", "text"),
        ("abcɭ", "garbled_text_layer"),
        ("123 ! \u0301", "page_no_text"),
        ("І\u0308 И\u0306 Ґґ єї \u0301", "text"),
        ("abcdeɭ\u0301" + "!123" * 20, "text"),
    ],
)
def test_pdf_letter_classifier_exact_boundary_normalisation_and_denominator(text, status):
    assert owned.page_text_status(text) == status


@pytest.mark.parametrize(
    "texts,expected,status",
    [
        (
            ["ɭɤɪɚʀғɧɫɶɤɨɸ", "ɭɤɪɚʀғɧɫɶɤɨɸ", "Readableword page."],
            ["garbled_text_layer", "garbled_text_layer", "text"],
            "extracted",
        ),
        (["ɭɤɪɚʀғɧɫɶɤɨɸ"], ["garbled_text_layer"], "skipped:garbled_text_layer"),
        (["123 !", "Readableword page."], ["page_no_text", "text"], "extracted"),
        (["abcdɭ", "abcɭ"], ["text", "garbled_text_layer"], "extracted"),
        ([owned.UKRAINIAN_LETTERS, "Clean English"], ["text", "text"], "extracted"),
    ],
)
def test_pdf_filters_garbled_pages_and_keeps_clean_pages(monkeypatch, texts, expected, status):
    # Stub only the PDF text layer: built-in PDF fonts cannot encode these scripts.
    monkeypatch.setattr(pymupdf.Page, "get_text", lambda page: texts[page.number])
    units, pages, actual = owned.extract(pdf(*("Placeholder" for _ in texts)), ".pdf")
    assert actual == status
    assert [p["status"] for p in pages] == expected
    assert units == [
        (i, text) for i, (text, state) in enumerate(zip(texts, expected, strict=True), 1) if state == "text"
    ]


def test_garbled_counts_are_separate_and_only_clean_text_is_stored(env, monkeypatch):
    db, root, _inventory, out = env
    texts = ["ɭɤɪɚʀғɧɫɶɤɨɸ", "ɭɤɪɚʀғɧɫɶɤɨɸ", "Readableword content.", "123!"]
    monkeypatch.setattr(pymupdf.Page, "get_text", lambda page: texts[page.number])
    (root / "fixture.pdf").write_bytes(pdf(*("Placeholder" for _ in texts)))
    request = args(env)
    report, code = owned.run(request)
    assert code == 0
    assert report["rows"][0]["chunks"] == 1
    assert report["rows"][0]["garbled_text_layer"] == 2
    assert report["rows"][0]["page_no_text"] == 1
    with sqlite3.connect(db) as conn:
        assert conn.execute("SELECT text FROM textbooks").fetchall() == [(texts[2],)]
    assert "ɭ" not in (out / "owned-synthetic-work.jsonl").read_text()
    request.check = True
    checked, code = owned.run(request)
    assert code == 0 and checked == report


@pytest.mark.parametrize("identities", [[], ["same", "same"], "legacy-book", [None], ["bad slug"]])
def test_invalid_source_files_rejected_with_typed_error(env, identities):
    with pytest.raises(owned.IngestError, match=r"^invalid_source_files$"):
        owned.run(args(env, [row(ingest="already_ingested", source_files=identities)]))


def test_singular_and_plural_identity_conflict_is_typed(env):
    with pytest.raises(owned.IngestError, match=r"^conflicting_source_identities$"):
        owned.run(args(env, [row(ingest="already_ingested", source_file="legacy", source_files=["legacy"])]))


def test_all_existing_identities_required_and_rights_preserved(env, monkeypatch):
    db, root, _inventory, out = env
    identities = [f"ulp-{i}-00-lesson-notes" for i in range(1, 7)]
    rights = owned.load_owned_rights()
    assert all(rights[identity]["rights"] == "owned_cite_only" for identity in identities)
    (root / "existing.pdf").write_bytes(pdf("Existingword notes."))
    request = args(env, [row(ingest="already_ingested", source_files=identities)])
    with sqlite3.connect(db) as conn:
        conn.executemany(
            "INSERT INTO textbooks(chunk_id,source_file,text) VALUES(?,?,?)",
            [(identity, identity, "Existing body") for identity in identities[:-1]],
        )
    report, code = owned.run(request)
    assert code == 1 and report["rows"][0]["status"] == "error:existing_source_missing"
    assert len(report["rows"][0]["sources"]) == 6
    with sqlite3.connect(db) as conn:
        conn.execute(
            "INSERT INTO textbooks(chunk_id,source_file,text) VALUES(?,?,?)",
            (identities[-1], identities[-1], "Existing body"),
        )
    before = db.read_bytes()
    for mode in ({}, {"check": True}):
        request.check = mode.get("check", False)
        report, code = owned.run(request)
        assert code == 0 and report["rows"][0]["status"] == "already_ingested"
        assert report["rows"][0]["chunks"] == 6
        assert [s["source_file"] for s in report["rows"][0]["sources"]] == identities
        assert all(s["rights"] == "owned_cite_only" for s in report["rows"][0]["sources"])
        assert db.read_bytes() == before and not out.exists()
    monkeypatch.setattr(owned, "load_owned_rights", lambda: {k: v for k, v in rights.items() if k != identities[-1]})
    assert owned.run(request)[0]["rows"][0]["status"] == "error:existing_source_rights_missing"


@pytest.mark.parametrize("state", ["unknown", "new_text_only"])
def test_unknown_ingest_modes_have_typed_error(env, state):
    report, code = owned.run(args(env, [row(ingest=state)], check=True))
    assert code == 1 and report["rows"][0]["status"] == "error:unknown_ingest"


@pytest.mark.parametrize("identity", [None, "", [], "bad slug"])
def test_invalid_singular_identity_has_typed_error(env, identity):
    with pytest.raises(owned.IngestError, match=r"^invalid_source_file$"):
        owned.run(args(env, [row(ingest="already_ingested", source_file=identity)]))


def test_unreadable_legacy_rights_have_sanitized_typed_error(env, monkeypatch):
    _db, root, _inventory, _out = env
    (root / "existing.pdf").write_bytes(pdf("Existingword notes."))

    def unreadable():
        raise ValueError("SENSITIVE_NAME_SENTINEL")

    monkeypatch.setattr(owned, "load_owned_rights", unreadable)
    report, code = owned.run(args(env, [row(ingest="already_ingested", source_file="legacy")]))
    assert code == 1 and report["rows"][0]["status"] == "error:owned_rights_unreadable"
    assert "SENSITIVE_NAME_SENTINEL" not in str(report)


@pytest.mark.parametrize("extension", [".pptx", ".epub"])
def test_packaged_text_filters_garbled_units_before_ingestion(monkeypatch, extension):
    monkeypatch.setattr(owned, "package_text", lambda payload, ext: [(1, "ɭɤɪɚʀғɧɫɶɤɨɸ"), (2, "Readableword unit.")])
    units, pages, status = owned.extract(b"synthetic", extension)
    assert status == "extracted" and units == [(2, "Readableword unit.")]
    assert [p["status"] for p in pages] == ["garbled_text_layer", "text"]


@pytest.mark.parametrize(
    "content",
    [
        '"First, quoted field",Anotherword\n"Second\nwrapped field",Moreword\n',
        'Firstword;Anotherword\n"Unescaped "quote"";Moreword\n',
        "Firstword\tAnotherword\nSecondword\tMoreword\n",
    ],
)
def test_csv_preserves_serialized_text_bom_quotes_and_separators(content):
    units, pages, status = owned.extract(("\ufeff" + content).encode(), ".csv")
    assert status == "extracted" and units == [(1, content)]
    assert [p["status"] for p in pages] == ["text"]


def test_csv_invalid_utf8_is_typed_corrupt():
    with pytest.raises(owned.IngestError, match=r"^corrupt$"):
        owned.extract(b"\xff", ".csv")


@pytest.mark.parametrize("extension", [".jpeg", ".jpg"])
def test_jpeg_pages_are_validated_and_accounted_as_ocr_residual(extension):
    image = pymupdf.Pixmap(pymupdf.csRGB, (0, 0, 4, 4), False)
    image.clear_with(255)
    units, pages, status = owned.extract(image.tobytes("jpeg"), extension)
    assert not units and status == "skipped:scanned_needs_ocr"
    assert pages == [{"page": 1, "status": "page_no_text", "characters": 0}]
    with pytest.raises(owned.IngestError, match=r"^corrupt$"):
        owned.extract(b"not a JPEG", extension)


def test_directory_inventory_accounts_csv_images_and_scanned_pdf(env):
    _db, root, _inventory, _out = env
    (root / "tree").mkdir()
    (root / "tree/notes.csv").write_text("Readableword,reference\nOtherword,notes\n")
    image = pymupdf.Pixmap(pymupdf.csRGB, (0, 0, 4, 4), False)
    image.clear_with(255)
    (root / "tree/image.jpeg").write_bytes(image.tobytes("jpeg"))
    (root / "tree/scanned.pdf").write_bytes(pdf(""))
    request = args(env, [row(files=["tree/**"])])
    report, code = owned.run(request)
    assert code == 0 and report["rows"][0]["chunks"] == 1
    assert report["rows"][0]["page_no_text"] == 2
    assert [f["status"] for f in report["rows"][0]["files"]] == [
        "skipped:scanned_needs_ocr",
        "extracted",
        "skipped:scanned_needs_ocr",
    ]
    request.check = True
    assert owned.run(request)[1] == 0


@pytest.mark.parametrize("packed", [False, True])
def test_exact_pdf_duplicates_across_disk_and_archive_ingest_once(env, packed):
    db, root, _inventory, _out = env
    payload = pdf("Distinctiveword shared page.", "Secondword shared page.")
    (root / "a.pdf").write_bytes(payload)
    (root / ("b.zip" if packed else "b.pdf")).write_bytes(
        archive([("member.pdf", payload)]) if packed else pdf("Distinctiveword shared page.", "Secondword shared page.")
    )
    report, code = owned.run(args(env, [row(files=["**/*"])]))
    assert code == 0 and report["rows"][0]["chunks"] == 2
    entries = report["rows"][0]["files"]
    assert entries[-1]["file_index"] == ("2.1" if packed else "2")
    assert entries[-1]["status"] == "duplicate_of:f1"
    assert entries[-1]["text_sha256"] == entries[0]["text_sha256"]
    with sqlite3.connect(db) as conn:
        assert conn.execute("SELECT count(*) FROM textbooks").fetchone()[0] == 2


@pytest.mark.parametrize(
    "texts,expected",
    [
        (["Prefixword middleword suffixword", "middleword"], ["extracted", "duplicate_of:f1"]),
        (
            ["Firstword middleword lastword", "Otherword middleword endword", "middleword"],
            ["extracted", "extracted", "duplicate_of:f1"],
        ),
        (["Alpha beta gamma", "gamma delta epsilon"], ["extracted", "extracted"]),
        # Matching separated regions is a partial overlap, not a substring.
        (["Firstword sharedword insertedword Lastword", "Firstword sharedword Lastword"], ["extracted", "extracted"]),
        (["Caf\u00e9 word", "Cafe\u0301\n\t word"], ["extracted", "duplicate_of:f1"]),
        (["Earlierword middleword", "middleword", "Earlierword"], ["extracted", "duplicate_of:f1", "duplicate_of:f1"]),
        # A later larger file is retained; only earlier retained files donate.
        (["middleword", "Prefixword middleword suffixword"], ["extracted", "extracted"]),
        # Whole-file containment remains independent of page and chunk boundaries.
        (["Prefixword " + "a" * (owned.CHUNK_SIZE + 10) + " Suffixword", "a" * 40], ["extracted", "duplicate_of:f1"]),
    ],
)
def test_normalised_whole_file_containment_earliest_donor_and_partial_overlap(env, texts, expected):
    _db, root, _inventory, _out = env
    for i, text in enumerate(texts):
        (root / f"{i}.csv").write_text(text)
    units, accounting = owned.collect(row(), sorted(root.glob("*.csv")))
    assert [f["status"] for f in accounting] == expected
    assert len(units) == expected.count("extracted")


def test_empty_text_files_keep_no_text_accounting_and_are_not_deduplicated(env):
    _db, root, _inventory, _out = env
    for name in ("a.pdf", "b.pdf"):
        (root / name).write_bytes(pdf(""))
    units, accounting = owned.collect(row(), sorted(root.glob("*.pdf")))
    assert not units
    assert all(f["status"] == "skipped:scanned_needs_ocr" for f in accounting)
    assert all("text_sha256" not in f for f in accounting)
    assert all(f["pages"][0]["status"] == "page_no_text" for f in accounting)


@pytest.mark.parametrize("extension", [".pdf", ".pptx"])
@pytest.mark.parametrize(
    "decks,expected",
    [
        (
            [["Alpha slide", "Donor only slide", "Beta slide"], ["Beta slide", "Alpha slide"]],
            ["extracted", "duplicate_of:f1"],
        ),
        (
            [["Alpha slide", "Donor only slide", "Beta slide"], ["Beta slide", "Alpha slide", "New slide"]],
            ["extracted", "extracted"],
        ),
        (
            [["Alpha slide", "First donor only"], ["Beta slide", "Second donor only"], ["Beta slide", "Alpha slide"]],
            ["extracted", "extracted", "extracted"],
        ),
        (
            [["Alpha slide", "Beta slide"], ["", ""]],
            ["extracted", "skipped:scanned_needs_ocr"],
        ),
        (
            [["Alpha slide", "Donor only slide", "Beta slide"], ["Beta slide", "", "Alpha   slide"]],
            ["extracted", "duplicate_of:f1"],
        ),
        (
            [
                ["Alpha slide", "First donor only", "Beta slide"],
                ["Beta slide", "Second donor only", "Alpha slide"],
                ["Beta slide", "Alpha slide"],
            ],
            ["extracted", "extracted", "duplicate_of:f1"],
        ),
    ],
)
def test_unit_duplicates_require_one_earlier_retained_donor(env, extension, decks, expected):
    db, root, _inventory, _out = env
    if extension == ".pptx":
        expected = ["skipped:page_no_text" if s == "skipped:scanned_needs_ocr" else s for s in expected]
    for index, slides in enumerate(decks, 1):
        if extension == ".pdf":
            payload = pdf(*slides)
        else:
            payload = archive(
                [
                    (
                        "ppt/presentation.xml",
                        '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                        "<p:sldIdLst>"
                        + "".join(f'<p:sldId id="{255 + i}" r:id="s{i}"/>' for i in range(1, len(slides) + 1))
                        + "</p:sldIdLst></p:presentation>",
                    ),
                    (
                        "ppt/_rels/presentation.xml.rels",
                        "<Relationships>"
                        + "".join(
                            f'<Relationship Id="s{i}" Target="slides/slide{i}.xml"/>' for i in range(1, len(slides) + 1)
                        )
                        + "</Relationships>",
                    ),
                    *[
                        (
                            f"ppt/slides/slide{i}.xml",
                            '<slide xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                            f"<a:p><a:t>{text}</a:t></a:p></slide>",
                        )
                        for i, text in enumerate(slides, 1)
                    ],
                ]
            )
        (root / f"{index}{extension}").write_bytes(payload)
    report, code = owned.run(args(env, [row(files=["**/*"])]))
    assert code == 0
    entries = report["rows"][0]["files"]
    assert [f["status"] for f in entries] == expected
    kept = [
        (str(i), j)
        for i, slides in enumerate(decks, 1)
        if expected[i - 1] == "extracted"
        for j, text in enumerate(slides, 1)
        if text.strip()
    ]
    with sqlite3.connect(db) as conn:
        stored = {r[0] for r in conn.execute("SELECT chunk_id FROM textbooks")}
    assert stored == {f"owned-synthetic-work_f{i}_p{j:04d}_c0001" for i, j in kept}


def test_unit_duplicate_normalisation_uses_nfc_and_collapsed_whitespace(env, monkeypatch):
    _db, root, _inventory, _out = env
    extracted = iter(
        [
            [(1, "Caf\u00e9 first slide"), (2, "Donor only slide"), (3, "Last slide")],
            [(1, "Last\n\t slide"), (2, "Cafe\u0301  first slide")],
        ]
    )
    monkeypatch.setattr(owned, "extract", lambda *a, **kw: (next(extracted), [], "extracted"))
    for name in ("a.pptx", "b.pptx"):
        (root / name).write_bytes(b"synthetic extraction input")
    units, accounting = owned.collect(row(), sorted(root.glob("*.pptx")))
    assert [f["status"] for f in accounting] == ["extracted", "duplicate_of:f1"]
    assert len(units) == 3 and all(index == "1" for index, _page, _text in units)


@pytest.mark.parametrize(
    "policy,expected", [({}, "missing"), ({"owned-synthetic-work": {"rights": "private_permission"}}, "mismatch")]
)
@pytest.mark.parametrize("mode", ["ingest", "check", "dry_run"])
def test_new_source_requires_matching_rights_before_any_writes(env, monkeypatch, policy, expected, mode):
    db, root, _inventory, out = env
    (root / "a.pdf").write_bytes(pdf("Readableword content."))
    request = args(env, **({mode: True} if mode != "ingest" else {}))
    before = db.read_bytes()
    monkeypatch.setattr(owned, "load_owned_rights", lambda: policy)
    report, code = owned.run(request)
    assert code == 1 and report["rows"][0]["status"] == f"error:owned_source_rights_{expected}"
    assert db.read_bytes() == before and not out.exists()


def test_private_inventory_cannot_use_cite_only_rights(env):
    db, root, _inventory, out = env
    (root / "a.pdf").write_bytes(pdf("Readableword content."))
    before = db.read_bytes()
    report, code = owned.run(args(env, [row(rights="private_permission")]))
    assert code == 1 and report["rows"][0]["status"] == "error:owned_source_rights_mismatch"
    assert db.read_bytes() == before and not out.exists()


@pytest.mark.parametrize("change", ["removed", "changed", "unreadable"])
@pytest.mark.parametrize("mode", ["rerun", "check", "cached_load", "force"])
def test_rights_withdrawal_blocks_valid_cached_artifacts_without_writes(env, monkeypatch, change, mode):
    db, root, _inventory, out = env
    (root / "a.pdf").write_bytes(pdf("Readableword content."))
    request = args(env)
    assert owned.run(request)[1] == 0
    if mode == "cached_load":
        target = db.with_name("fresh.db")
        with sqlite3.connect(target) as conn:
            conn.executescript(SCHEMA)
        request.db = target
    request.check = mode == "check"
    request.force = mode == "force"
    before = request.db.read_bytes(), {p.name: (p.read_bytes(), p.stat().st_mtime_ns) for p in out.iterdir()}
    if change == "unreadable":

        def unreadable():
            raise ValueError("SENSITIVE_SENTINEL")

        monkeypatch.setattr(owned, "load_owned_rights", unreadable)
        error = "owned_rights_unreadable"
    else:
        monkeypatch.setattr(
            owned,
            "load_owned_rights",
            lambda: {} if change == "removed" else {"owned-synthetic-work": {"rights": "private_permission"}},
        )
        error = f"owned_source_rights_{'missing' if change == 'removed' else 'mismatch'}"
    report, code = owned.run(request)
    assert code == 1 and report["rows"][0]["status"] == f"error:{error}"
    assert before == (request.db.read_bytes(), {p.name: (p.read_bytes(), p.stat().st_mtime_ns) for p in out.iterdir()})


def test_disk_metadata_files_are_ignored_and_check_passes(env):
    _db, root, _inventory, _out = env
    (root / "a.pdf").write_bytes(pdf("Readableword content."))
    for name in (".DS_Store", "._x.pdf"):
        (root / name).write_bytes(b"not valid documents")
    request = args(env, [row(files=["**/*"])])
    report, code = owned.run(request)
    assert code == 0 and report["rows"][0]["chunks"] == 1
    assert [f["status"] for f in report["rows"][0]["files"]] == [
        "skipped:ignored_metadata",
        "skipped:ignored_metadata",
        "extracted",
    ]
    request.check = True
    assert owned.run(request)[1] == 0


@pytest.mark.parametrize("packed", [False, True])
@pytest.mark.parametrize("user_password", ["", "synthetic-user"])
def test_owner_password_is_accounted_but_user_password_is_refused(env, packed, user_password):
    _db, root, _inventory, _out = env
    with pymupdf.open(stream=pdf("Readableword restricted content."), filetype="pdf") as doc:
        payload = doc.tobytes(
            encryption=pymupdf.PDF_ENCRYPT_AES_256,
            owner_pw="synthetic-owner",
            user_pw=user_password,
            permissions=pymupdf.PDF_PERM_PRINT,
        )
    (root / ("a.zip" if packed else "a.pdf")).write_bytes(archive([("member.pdf", payload)]) if packed else payload)
    request = args(env, [row(files=["**/*"])])
    report, code = owned.run(request)
    entry = report["rows"][0]["files"][-1]
    if user_password:
        assert code == 1 and entry["status"] == "error:encrypted"
        assert "owner_restricted" not in entry
    else:
        assert code == 0 and entry["status"] == "extracted"
        assert entry["owner_restricted"] is True
        assert report["rows"][0]["chunks"] == 1
        request.check = True
        assert owned.run(request)[1] == 0
